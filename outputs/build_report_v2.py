from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import copy

# カラー定義
C_BG       = RGBColor(0xFF, 0xFF, 0xFF)
C_ACCENT   = RGBColor(0x00, 0x72, 0xC6)   # シナプスブルー系
C_ACCENT2  = RGBColor(0x00, 0xB0, 0x50)   # グリーン（改善）
C_WARN     = RGBColor(0xFF, 0x6B, 0x00)   # オレンジ（課題）
C_DARK     = RGBColor(0x1F, 0x1F, 0x1F)
C_GRAY     = RGBColor(0x76, 0x76, 0x76)
C_LIGHT    = RGBColor(0xF2, 0xF6, 0xFB)
C_WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
C_RED      = RGBColor(0xC0, 0x00, 0x00)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]  # 完全ブランク


def add_slide():
    return prs.slides.add_slide(blank_layout)


def rect(slide, l, t, w, h, fill=None, line=None, line_w=Pt(0)):
    from pptx.util import Pt
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.width = line_w
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def txbox(slide, text, l, t, w, h, size=14, bold=False, color=C_DARK,
          align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def header(slide, title, subtitle=""):
    rect(slide, 0, 0, 13.33, 1.1, fill=C_ACCENT)
    txbox(slide, title, 0.4, 0.15, 10, 0.55, size=24, bold=True, color=C_WHITE)
    if subtitle:
        txbox(slide, subtitle, 0.4, 0.68, 10, 0.35, size=13, color=RGBColor(0xCC,0xE4,0xFF))


def kpi_card(slide, label, value, unit="", l=0, t=0, w=2.8, h=1.4,
             val_color=C_ACCENT, bg=C_LIGHT):
    rect(slide, l, t, w, h, fill=bg, line=RGBColor(0xCC,0xD9,0xEA), line_w=Pt(1))
    txbox(slide, label, l+0.15, t+0.1, w-0.2, 0.35, size=11, color=C_GRAY)
    txbox(slide, value, l+0.15, t+0.4, w-0.2, 0.7, size=28, bold=True, color=val_color)
    if unit:
        txbox(slide, unit, l+0.15, t+1.05, w-0.2, 0.28, size=10, color=C_GRAY)


def table_block(slide, headers, rows, l, t, w, h,
                header_fill=C_ACCENT, header_color=C_WHITE,
                row_fills=None):
    col_n = len(headers)
    col_w = w / col_n
    row_n = len(rows) + 1
    row_h = h / row_n

    # ヘッダー行
    for ci, hd in enumerate(headers):
        rect(slide, l + ci*col_w, t, col_w, row_h, fill=header_fill)
        txbox(slide, hd, l+ci*col_w+0.07, t+0.04, col_w-0.1, row_h-0.05,
              size=10, bold=True, color=header_color, align=PP_ALIGN.CENTER)

    for ri, row in enumerate(rows):
        bg = RGBColor(0xF2,0xF6,0xFB) if ri % 2 == 0 else C_WHITE
        if row_fills and ri < len(row_fills) and row_fills[ri]:
            bg = row_fills[ri]
        for ci, cell in enumerate(row):
            rect(slide, l+ci*col_w, t+(ri+1)*row_h, col_w, row_h,
                 fill=bg, line=RGBColor(0xCC,0xD9,0xEA), line_w=Pt(0.5))
            is_bold = str(cell).startswith("**") or str(cell).startswith("✅") or str(cell).startswith("🔴")
            clean = str(cell).replace("**","")
            txbox(slide, clean, l+ci*col_w+0.07, t+(ri+1)*row_h+0.04,
                  col_w-0.1, row_h-0.06, size=10, bold=is_bold,
                  color=C_DARK, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────
# スライド 1：表紙
# ─────────────────────────────────────────
s = add_slide()
rect(s, 0, 0, 13.33, 7.5, fill=C_ACCENT)
rect(s, 0, 4.5, 13.33, 3.0, fill=RGBColor(0x00,0x55,0x99))
txbox(s, "SNS CM制作・広告運用", 1.5, 1.5, 10, 0.8,
      size=20, color=RGBColor(0xCC,0xE4,0xFF), align=PP_ALIGN.CENTER)
txbox(s, "活動報告", 1.5, 2.2, 10, 1.4,
      size=48, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
txbox(s, "2026年6月5日〜6月18日（2週間）", 1.5, 3.7, 10, 0.6,
      size=18, color=RGBColor(0xCC,0xE4,0xFF), align=PP_ALIGN.CENTER)
txbox(s, "作成日：2026-06-18　　株式会社シナプス 営業課", 1.5, 5.8, 10, 0.5,
      size=13, color=RGBColor(0xAA,0xCC,0xFF), align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────
# スライド 2：KPIサマリー
# ─────────────────────────────────────────
s = add_slide()
header(s, "KPI サマリー", "2026/06/05〜06/17　比較対象：前期（5/27〜6/2）")

kpi_card(s, "総消化金額", "¥8,537", "前期 ¥3,329",   l=0.4,  t=1.3, val_color=C_DARK)
kpi_card(s, "LPV件数",   "74件",   "前期 29件 ▲155%", l=3.4,  t=1.3, val_color=C_ACCENT2)
kpi_card(s, "LPV単価",   "¥58",    "前期 ¥115 ▼50%改善",l=6.4, t=1.3, val_color=C_ACCENT2)
kpi_card(s, "フォームリード", "1件", "前期 0件 ★初獲得", l=9.4, t=1.3, val_color=C_ACCENT2)

kpi_card(s, "フック率（LPV）", "32%", "前作 31% 過去最高", l=0.4, t=3.0, val_color=C_ACCENT)
kpi_card(s, "継続率（LPV）",  "3.84%","前作 6% ⚠課題",   l=3.4, t=3.0, val_color=C_WARN)
kpi_card(s, "継続率（フォーム）","13.29%","全パターン最高", l=6.4,t=3.0, val_color=C_ACCENT2)
kpi_card(s, "リード単価",    "¥4,256","フォームキャンペーン", l=9.4, t=3.0, val_color=C_DARK)

txbox(s, "✅ フォーム修正（Japanese locale化）の効果が出始めた可能性　　⚠ 中盤離脱の改善が次の最優先課題",
      0.4, 5.0, 12.5, 0.5, size=12, color=C_GRAY)


# ─────────────────────────────────────────
# スライド 3：A/Bテスト結果
# ─────────────────────────────────────────
s = add_slide()
header(s, "A/Bテスト：LP経由 vs Meta直接フォーム")

# 左パネル
rect(s, 0.4, 1.2, 5.8, 5.5, fill=C_LIGHT, line=RGBColor(0xCC,0xD9,0xEA), line_w=Pt(1))
txbox(s, "パターンA　LP経由", 0.5, 1.25, 5.6, 0.5, size=14, bold=True, color=C_ACCENT)
txbox(s, "広告　→　LP　→　フォーム", 0.6, 1.75, 5.4, 0.4, size=12, color=C_GRAY)

for label, val in [("消化金額","¥4,281"),("LPV件数","74件"),("リード獲得","0件"),("リード単価","—")]:
    pass

table_block(s,
    ["指標","数値"],
    [["消化金額","¥4,281"],
     ["LPV件数","74件"],
     ["リード獲得","0件（メールなし）"],
     ["リード単価","—"]],
    l=0.5, t=2.2, w=5.6, h=2.8)

txbox(s, "LPに74件誘導するも問い合わせゼロ", 0.5, 5.1, 5.6, 0.4, size=11, color=C_WARN, bold=True)

# 右パネル
rect(s, 7.0, 1.2, 5.8, 5.5, fill=RGBColor(0xE8,0xF5,0xEE),
     line=C_ACCENT2, line_w=Pt(2))
txbox(s, "パターンB　Meta直接フォーム ★", 7.1, 1.25, 5.6, 0.5, size=14, bold=True, color=C_ACCENT2)
txbox(s, "広告　→　Metaフォーム（LP不要）", 7.2, 1.75, 5.4, 0.4, size=12, color=C_GRAY)

table_block(s,
    ["指標","数値"],
    [["消化金額","¥4,256"],
     ["リード獲得","★ 1件"],
     ["リード単価","¥4,256"],
     ["継続率","13.29%（全体最高）"]],
    l=7.1, t=2.2, w=5.6, h=2.8,
    header_fill=C_ACCENT2)

txbox(s, "初リード獲得！直接フォームが優位", 7.1, 5.1, 5.6, 0.4, size=11, color=C_ACCENT2, bold=True)

txbox(s, "※ サンプルが少ないため引き続き観察が必要。2〜3週間データを積んだ後に予算配分を判断する。",
      0.4, 6.85, 12.5, 0.45, size=10, color=C_GRAY)


# ─────────────────────────────────────────
# スライド 4：動画パフォーマンス比較
# ─────────────────────────────────────────
s = add_slide()
header(s, "動画パフォーマンス比較（全履歴）")

table_block(s,
    ["動画クリエイティブ","フック率","継続率","備考"],
    [["LPVキャンペーン（駆けつけ訴求）","32%\n▲過去最高","3.84%\n⚠最低水準","フックした後の中盤離脱が課題"],
     ["フォームキャンペーン（桜島俯瞰）","24%","13.29%\n★全体最高","中盤〜後半に引きつける力がある"],
     ["前作パターンA（参考）","31%","6%","基準値"],
     ["前作パターンB（参考）","24%","6%","基準値"]],
    l=0.5, t=1.3, w=12.3, h=3.6,
    row_fills=[None, RGBColor(0xE8,0xF5,0xEE), None, None])

rect(s, 0.5, 5.1, 12.3, 1.8, fill=RGBColor(0xFF,0xF3,0xE8),
     line=C_WARN, line_w=Pt(1))
txbox(s, "⚠ 課題：フックした後の離脱", 0.7, 5.15, 6, 0.4, size=13, bold=True, color=C_WARN)
txbox(s,
    "フック率32%で引きつけた視聴者の大半が4秒以降で離脱。\n"
    "乗り換えCM v2では中盤（シーン②③）の構成強化が最重要ポイント。",
    0.7, 5.55, 11.8, 0.9, size=11, color=C_DARK)


# ─────────────────────────────────────────
# スライド 5：実施施策と結果
# ─────────────────────────────────────────
s = add_slide()
header(s, "今期の実施施策と結果")

table_block(s,
    ["施策","実施時期","結果・ステータス"],
    [["フォームのJapanese locale化","6月上旬","✅ フォームリード初獲得に寄与した可能性"],
     ["確認ステップ廃止（大量用フォームに変更）","6月上旬","✅ 申込ハードル低下"],
     ["パターンBの広告フォームを差し替え","6月上旬","✅ リード1件獲得"],
     ["AIエージェントパイプライン構築","6月上旬","✅ CM制作フローを完全整備"],
     ["乗り換えCM v2 ブリーフ・プロンプト作成","6月中旬","🔄 動画化・編集を進行中"]],
    l=0.5, t=1.3, w=12.3, h=4.5)


# ─────────────────────────────────────────
# スライド 6：CM制作進捗
# ─────────────────────────────────────────
s = add_slide()
header(s, "CM制作進捗｜乗り換えCM v2（巻き込まれ感からの解放）")

table_block(s,
    ["工程","状況","備考"],
    [["ブリーフ","✅ 完成","briefs/brief_norikae_v2.md"],
     ["コンセプト・台本・絵コンテ","✅ v3確定","前作から継承・改良"],
     ["静止画生成プロンプト（全シーン）","✅ 完成","①②B③A③B④A④A-2"],
     ["Flow I2Vプロンプト（全シーン）","✅ 完成","Veo 3.1 Fast / Veo標準（④Aのみ）"],
     ["動画化（Veo 3.1 Fast）","🔄 進行中","4秒×各シーン・1発生成ルール"],
     ["Canva合成素材","⏳ 待機中","②A・④A-2のUI文言"],
     ["Premiere編集・書き出し","⏳ 待機中","9:16 / 1080×1920 / 30秒"]],
    l=0.5, t=1.3, w=12.3, h=4.5)

txbox(s, "v2の主な改善：③B駆けつけを車両到着案に固定　／　③ナレーションを20文字×2文に制限　／　④A-2（資料請求シーン）を新規追加",
      0.5, 6.1, 12.3, 0.5, size=10, color=C_GRAY)


# ─────────────────────────────────────────
# スライド 7：次のアクション
# ─────────────────────────────────────────
s = add_slide()
header(s, "次のアクション")

actions = [
    ("🔴 最優先", "乗り換えCM v2　動画素材完成・Premiere編集・書き出し",         C_RED),
    ("🔴 最優先", "v2をLPVキャンペーンのパターンB（フック率24%）と差し替え配信",   C_RED),
    ("🟡 次点",   "フォームリードの継続観察（1件が偶発か傾向かを判断）",            C_WARN),
    ("🟡 次点",   "v2配信開始後3〜4日でフック率・継続率を前作と比較",              C_WARN),
    ("🟢 中期",   "継続率改善を確認後、予算増額・パターンA差し替えを判断",          C_ACCENT2),
]

for i, (priority, task, color) in enumerate(actions):
    y = 1.3 + i * 1.0
    rect(s, 0.4, y, 1.5, 0.75, fill=color)
    txbox(s, priority, 0.42, y+0.08, 1.45, 0.55, size=12, bold=True,
          color=C_WHITE, align=PP_ALIGN.CENTER)
    rect(s, 2.1, y, 10.8, 0.75, fill=C_LIGHT,
         line=RGBColor(0xCC,0xD9,0xEA), line_w=Pt(1))
    txbox(s, task, 2.25, y+0.15, 10.5, 0.5, size=13, color=C_DARK)


# 保存
out = r"c:\Users\fukuyama\Claude Code\SNS\outputs\活動報告_2026-06-05_2026-06-18.pptx"
prs.save(out)
print(f"saved: {out}")
